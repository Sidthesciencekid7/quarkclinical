"""Focused tests for the budget Word (.docx) export
(routers/budgets.py::export_budget_docx, services/budget_docx_export.py).

Kept in its own file (rather than appended to tests/test_budgets.py) so it can
be run in isolation without racing the parallel work landing in that file, per
the task instructions. Reuses tests/test_budgets.py's local helpers via a
thin import rather than duplicating them.
"""
from __future__ import annotations

import io
import uuid

import docx
import pytest

from models import User, UserRole
from security import hash_password
from tests.test_budgets import _as, _create_budget

_DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


async def _make_outsider(db) -> User:
    outsider = User(
        id=uuid.uuid4(),
        org_id="other-org-docx",
        email="outsider-docx@other.example",
        role=UserRole.generator,
        hashed_password=hash_password("irrelevant"),
    )
    db.add(outsider)
    await db.commit()
    await db.refresh(outsider)
    return outsider


async def test_export_docx_returns_a_valid_populated_document(api_client, seeded):
    created = await _create_budget(
        api_client, seeded.user, name="Metformin BE Budget", drug_inn="metformin",
    )
    budget_id = created["id"]

    resp = await api_client.get(
        f"/api/budgets/{budget_id}/export/docx", headers=_as(seeded.user),
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == _DOCX_MEDIA_TYPE
    assert "attachment" in resp.headers["content-disposition"]

    body = resp.content
    assert body[:4] == b"PK\x03\x04"
    assert len(body) > 10_000

    # Fetch the rollup so the test can assert against the SAME numbers the
    # document renders, rather than a hardcoded figure.
    rollup_resp = await api_client.get(
        f"/api/budgets/{budget_id}/rollup", headers=_as(seeded.user),
    )
    rollup = rollup_resp.json()
    grand_total_str = f"{rollup['grand_total']:,.2f}"

    document = docx.Document(io.BytesIO(body))
    full_text = "\n".join(p.text for p in document.paragraphs)
    full_text += "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )

    assert "Summary Study Budget" in full_text
    assert "Itemized Study Budget" in full_text
    assert "GRAND TOTAL" in full_text
    assert grand_total_str in full_text


async def test_export_docx_cross_org_is_forbidden_and_unknown_is_404(api_client, seeded):
    outsider = await _make_outsider(seeded.db)
    created = await _create_budget(api_client, seeded.user, drug_inn="metformin")
    budget_id = created["id"]

    resp = await api_client.get(
        f"/api/budgets/{budget_id}/export/docx", headers=_as(outsider),
    )
    assert resp.status_code == 403

    resp = await api_client.get(
        f"/api/budgets/{uuid.uuid4()}/export/docx", headers=_as(seeded.user),
    )
    assert resp.status_code == 404
