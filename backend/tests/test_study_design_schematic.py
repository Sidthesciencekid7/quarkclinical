"""Tests for the Figure 1 study-design schematic (DOCX + PDF, routers/export.py).

A parallel-group study runs two treatment arms concurrently; a crossover study
moves the same subjects through periods sequentially. The schematic is
supposed to show that difference: a parallel design's two arms stay in their
own visually separate lane through every stage they are actually in (not just
a single box naming both arms, and not a momentary split that immediately
re-merges), while a crossover design - which never has two arms running at
once - stays a single row.

Rendering is deterministic (no AI, no randomness), so the exact shape of the
output is pinned directly.
"""
from __future__ import annotations

import io

import pytest
from docx import Document
from reportlab.graphics.shapes import Line, Rect
from reportlab.platypus import SimpleDocTemplate
from reportlab.lib.pagesizes import letter

from routers.export import (
    _add_study_design_schematic_docx,
    _add_study_design_schematic_pdf,
    _design_schematic_stages,
    _pdf_styles,
)


class _FakeProtocol:
    def __init__(self, sample_size: dict | None):
        self.population_config = {"_sample_size": sample_size} if sample_size else {}


_PARALLEL_REC = {
    "design": "Parallel-Group", "design_key": "parallel",
    "n_enrolled": 94, "n_per_group": 47,
}
_CROSSOVER_REC = {
    "design": "2x2 Crossover", "design_key": "crossover",
    "n_enrolled": 32, "n_per_group": None,
}


# ---------------------------------------------------------------------------
# _design_schematic_stages — the shared data both renderers draw from
# ---------------------------------------------------------------------------


def test_parallel_stages_branch_at_arms_and_assessment_only():
    stages = _design_schematic_stages(_PARALLEL_REC)
    branch_flags = [bool(s["arms"]) for s in stages]
    # Screening, Randomize shared; arm + assessment branched; EOS shared again.
    assert branch_flags == [False, False, True, True, False]
    assert stages[2]["arms"] == ["Test Product (n=47)", "Reference (n=47)"]
    assert stages[3]["arms"][0] == stages[3]["arms"][1]  # same assessment, two lanes
    assert stages[0]["label"] == "Screening (Day -28 to -1)"
    assert stages[-1]["label"] == "End of Study (EOS)"


def test_crossover_stages_never_branch():
    stages = _design_schematic_stages(_CROSSOVER_REC)
    assert all(s["arms"] is None for s in stages)
    assert [s["label"] for s in stages] == [
        "Screening",
        "Randomize sequence (N=32)",
        "Period 1 (Test or Reference)",
        "Washout (>=5-7 x t1/2)",
        "Period 2 (crossover)",
        "End of Study (EOS)",
    ]


def test_stages_fall_back_to_generic_n_when_counts_missing():
    stages = _design_schematic_stages({"design_key": "parallel"})
    assert stages[1]["label"] == "Randomize 1:1 (N)"
    assert stages[2]["arms"] == ["Test Product (n)", "Reference (n)"]


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def test_docx_parallel_schematic_uses_two_rows_with_arms_in_separate_rows():
    doc = Document()
    _add_study_design_schematic_docx(doc, _FakeProtocol(_PARALLEL_REC))

    tbl = doc.tables[0]
    assert len(tbl.rows) == 2

    row0 = [c.text for c in tbl.rows[0].cells]
    row1 = [c.text for c in tbl.rows[1].cells]

    # Shared stages (Screening, Randomize, EOS) are the SAME merged cell, so
    # they read identically in both rows.
    for shared_text in ("Screening (Day -28 to -1)", "Randomize 1:1 (N=94)", "End of Study (EOS)"):
        assert shared_text in row0
        assert shared_text in row1

    # The two arms must be in DIFFERENT rows, not the same cell.
    assert "Test Product (n=47)" in row0
    assert "Test Product (n=47)" not in row1
    assert "Reference (n=47)" in row1
    assert "Reference (n=47)" not in row0
    # Neither row may contain both arm names slash-joined in one cell - that
    # was the old, fixed behaviour.
    assert not any("Test Product" in c and "Reference" in c for c in row0 + row1)


def test_docx_crossover_schematic_stays_a_single_row():
    doc = Document()
    _add_study_design_schematic_docx(doc, _FakeProtocol(_CROSSOVER_REC))

    tbl = doc.tables[0]
    assert len(tbl.rows) == 1
    texts = [c.text for c in tbl.rows[0].cells]
    assert "Period 1 (Test or Reference)" in texts
    assert "Period 2 (crossover)" in texts


def test_docx_schematic_omitted_when_no_sample_size_recommendation():
    doc = Document()
    _add_study_design_schematic_docx(doc, _FakeProtocol(None))
    assert len(doc.tables) == 0


def test_docx_parallel_schematic_saves_without_error():
    doc = Document()
    _add_study_design_schematic_docx(doc, _FakeProtocol(_PARALLEL_REC))
    buf = io.BytesIO()
    doc.save(buf)
    assert buf.tell() > 0


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _drawing_from_story(story):
    drawings = [el for el in story if hasattr(el, "contents")]
    assert len(drawings) == 1
    return drawings[0]


def test_pdf_parallel_schematic_draws_two_boxes_per_branched_stage():
    styles = _pdf_styles()
    story = []
    _add_study_design_schematic_pdf(story, _FakeProtocol(_PARALLEL_REC), styles, 3)
    drawing = _drawing_from_story(story)

    rects = [c for c in drawing.contents if isinstance(c, Rect)]
    # 3 shared stages (1 box each) + 2 branched stages (2 boxes each) = 7.
    assert len(rects) == 7

    lines = [c for c in drawing.contents if isinstance(c, Line)]
    # shared->shared (1) + branch (2x2 diagonal+horizontal segments) +
    # branched->branched (2 parallel arrows) + converge (2x2 segments) = 11.
    assert len(lines) == 11


def test_pdf_crossover_schematic_stays_unbranched():
    styles = _pdf_styles()
    story = []
    _add_study_design_schematic_pdf(story, _FakeProtocol(_CROSSOVER_REC), styles, 3)
    drawing = _drawing_from_story(story)

    rects = [c for c in drawing.contents if isinstance(c, Rect)]
    assert len(rects) == 6  # one box per stage, none split

    lines = [c for c in drawing.contents if isinstance(c, Line)]
    assert len(lines) == 5  # one straight arrow per gap, 5 gaps


def test_pdf_parallel_arm_boxes_sit_in_two_distinct_lanes():
    """The two arm boxes for a branched stage must be at different y-centres
    (separate visual lanes) - not stacked at the same height, and not the
    same box reused."""
    styles = _pdf_styles()
    story = []
    _add_study_design_schematic_pdf(story, _FakeProtocol(_PARALLEL_REC), styles, 3)
    drawing = _drawing_from_story(story)
    rects = [c for c in drawing.contents if isinstance(c, Rect)]

    # Group by x (left edge, rounded) to find the branched columns (2 boxes
    # sharing an x but at different y).
    from collections import defaultdict
    by_x = defaultdict(list)
    for r in rects:
        by_x[round(r.x, 1)].append(r.y)
    branched_columns = [ys for ys in by_x.values() if len(ys) == 2]
    assert len(branched_columns) == 2  # arm stage + assessment stage
    for ys in branched_columns:
        assert ys[0] != ys[1]


def test_pdf_schematic_falls_back_to_generic_crossover_shape_when_no_sample_size():
    """Pre-existing behaviour, unchanged by this fix: unlike the DOCX path
    (which skips the figure entirely - see
    test_docx_schematic_omitted_when_no_sample_size_recommendation), the PDF
    path has no `_sample_size` record to key off, defaults design_key to
    "crossover", and draws a generic, unbranched, placeholder-labelled
    schematic rather than omitting the figure."""
    styles = _pdf_styles()
    story = []
    _add_study_design_schematic_pdf(story, _FakeProtocol(None), styles, 3)
    assert story != []
    drawing = _drawing_from_story(story)
    rects = [c for c in drawing.contents if isinstance(c, Rect)]
    assert len(rects) == 6  # crossover shape: one box per stage, unbranched


def test_pdf_schematic_omitted_for_unknown_design_key():
    styles = _pdf_styles()
    story = []
    rec = {"design": "Unusual Design", "design_key": "something_else"}
    _add_study_design_schematic_pdf(story, _FakeProtocol(rec), styles, 3)
    assert story == []


def test_pdf_parallel_schematic_builds_without_error():
    styles = _pdf_styles()
    story = []
    _add_study_design_schematic_pdf(story, _FakeProtocol(_PARALLEL_REC), styles, 3)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    doc.build(story)
    assert buf.tell() > 0
