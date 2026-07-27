"""Integration tests for the Clinical Study Report (CSR) endpoints.

Covers the full backend flow on the validated engines: create a CSR shell from a
completed protocol, upload a subject-level results file, and get back the LOCKED
deterministic analysis (per-metric GMR + 90% CI + verdict). Also covers template
download, strict upload validation, and org isolation.

External calls (PubMed/DailyMed/Anthropic) are mocked exactly as in
test_pipeline.py / test_risk.py; the analysis itself is pure Python (the BE
engine's optional R cross-check is a subprocess and never hits the network).
"""

from __future__ import annotations

import csv
import io
import uuid

from models import User, UserRole
from security import hash_password
from tests.conftest import register_pubmed_mocks

_STUDY_PAYLOAD_BASE = {
    "study_type": "BE",
    "population_config": {"dose_strength": "850 mg", "formulation": "tablet"},
    "org_id": "test-org",
    "guideline_version": "FDA-BE-2014",
}


def _protocol_payload(drug_id: str, user_id: str) -> dict:
    return {**_STUDY_PAYLOAD_BASE, "drug_id": drug_id, "created_by": user_id}


async def _create_protocol(api_client, seeded, httpx_mock) -> str:
    register_pubmed_mocks(httpx_mock)
    resp = await api_client.post(
        "/api/protocols/",
        json=_protocol_payload(str(seeded.drug.id), str(seeded.user.id)),
        headers={"X-User-Id": str(seeded.user.id)},
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


async def _create_study_report(api_client, seeded, protocol_id, *, input_kind="params") -> str:
    resp = await api_client.post(
        f"/api/protocols/{protocol_id}/study-report",
        json={"input_kind": input_kind},
        headers={"X-User-Id": str(seeded.user.id)},
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["input_kind"] == input_kind
    assert body["results"] == {}
    return body["id"]


# A near-bioequivalent 2x2 crossover (T/R ratios ~1.01, low variability → the
# 90% CI sits well inside 80-125%). auc columns scale with cmax.
_PAIRS = [
    ("S01", "TR", 101, 100), ("S02", "RT", 103, 100), ("S03", "TR", 99, 100),
    ("S04", "RT", 102, 100), ("S05", "TR", 101, 99), ("S06", "RT", 104, 101),
    ("S07", "TR", 98, 100), ("S08", "RT", 102, 101), ("S09", "TR", 100, 99),
    ("S10", "RT", 103, 101),
]


def _params_csv(pairs=_PAIRS) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["subject", "sequence", "period", "treatment", "cmax", "auc0t", "auc0inf"])
    for subj, seq, tc, rc in pairs:
        if seq == "TR":
            plan = [("1", "T", tc), ("2", "R", rc)]
        else:
            plan = [("1", "R", rc), ("2", "T", tc)]
        for period, trt, cm in plan:
            w.writerow([subj, seq, period, trt, cm, round(cm * 5, 3), round(cm * 5.5, 3)])
    return buf.getvalue().encode("utf-8")


def _upload(api_client, seeded, sr_id, data: bytes, filename="results.csv"):
    return api_client.post(
        f"/api/study-reports/{sr_id}/dataset",
        files={"file": (filename, data, "text/csv")},
        headers={"X-User-Id": str(seeded.user.id)},
    )


# ---------------------------------------------------------------------------
# Happy path: create → upload → deterministic locked results
# ---------------------------------------------------------------------------


async def test_full_csr_flow_params(api_client, seeded, mock_anthropic, httpx_mock):
    protocol_id = await _create_protocol(api_client, seeded, httpx_mock)
    sr_id = await _create_study_report(api_client, seeded, protocol_id)

    resp = await _upload(api_client, seeded, sr_id, _params_csv())
    assert resp.status_code == 200, resp.text
    body = resp.json()

    # The dataset was analysed and the numbers locked in.
    results = body["results"]
    assert results["input_kind"] == "params"
    assert results["verdict"] == "bioequivalent"
    assert results["disposition"]["n_subjects"] == 10

    be = results["be"]
    for metric in ("cmax", "auc0t", "auc0inf"):
        r = be[metric]
        assert r["within_be_limits"] is True, (metric, r)
        # near-identical T/R → point estimate close to 100%
        assert 95.0 < r["point_estimate_pct"] < 115.0
        assert r["ci_lower_pct"] < r["point_estimate_pct"] < r["ci_upper_pct"]
        assert r["df"] == 8  # 20 obs - (intercept + treatment + period + 9 subject dummies)

    # A checksum + non-identifying summary are retained; raw rows are not stored.
    assert len(body["dataset_checksum"]) == 64
    assert body["dataset_summary"]["n_subjects"] == 10


async def test_get_and_list_study_report(api_client, seeded, mock_anthropic, httpx_mock):
    protocol_id = await _create_protocol(api_client, seeded, httpx_mock)
    sr_id = await _create_study_report(api_client, seeded, protocol_id)
    await _upload(api_client, seeded, sr_id, _params_csv())

    got = await api_client.get(
        f"/api/study-reports/{sr_id}", headers={"X-User-Id": str(seeded.user.id)}
    )
    assert got.status_code == 200
    assert got.json()["results"]["verdict"] == "bioequivalent"

    listing = await api_client.get(
        f"/api/protocols/{protocol_id}/study-reports",
        headers={"X-User-Id": str(seeded.user.id)},
    )
    assert listing.status_code == 200
    rows = listing.json()
    assert len(rows) == 1
    assert rows[0]["id"] == sr_id
    assert rows[0]["has_results"] is True
    assert rows[0]["verdict"] == "bioequivalent"


# ---------------------------------------------------------------------------
# Templates
# ---------------------------------------------------------------------------


async def test_template_download(api_client, seeded):
    for kind in ("nca", "params", "ae"):
        resp = await api_client.get(
            f"/api/study-reports/template?kind={kind}&fmt=csv",
            headers={"X-User-Id": str(seeded.user.id)},
        )
        assert resp.status_code == 200, resp.text
        assert b"subject" in resp.content

    xlsx = await api_client.get(
        "/api/study-reports/template?kind=nca&fmt=xlsx",
        headers={"X-User-Id": str(seeded.user.id)},
    )
    assert xlsx.status_code == 200
    assert xlsx.content[:2] == b"PK"  # xlsx is a zip

    bad = await api_client.get(
        "/api/study-reports/template?kind=nope&fmt=csv",
        headers={"X-User-Id": str(seeded.user.id)},
    )
    assert bad.status_code == 422


# ---------------------------------------------------------------------------
# Strict validation + guards
# ---------------------------------------------------------------------------


async def test_bad_dataset_returns_422_with_errors(api_client, seeded, mock_anthropic, httpx_mock):
    protocol_id = await _create_protocol(api_client, seeded, httpx_mock)
    sr_id = await _create_study_report(api_client, seeded, protocol_id)

    bad = (
        "subject,sequence,period,treatment,cmax,auc0t,auc0inf\n"
        "S01,TR,1,X,100,500,550\n"          # treatment X
        "S02,TR,notanint,T,90,450,500\n"    # period not int
    ).encode()
    resp = await _upload(api_client, seeded, sr_id, bad)
    assert resp.status_code == 422, resp.text
    detail = resp.json()["detail"]
    assert "errors" in detail and len(detail["errors"]) >= 2


async def test_empty_file_rejected(api_client, seeded, mock_anthropic, httpx_mock):
    protocol_id = await _create_protocol(api_client, seeded, httpx_mock)
    sr_id = await _create_study_report(api_client, seeded, protocol_id)
    resp = await _upload(api_client, seeded, sr_id, b"")
    assert resp.status_code == 422


async def test_unknown_study_report_404(api_client, seeded):
    resp = await api_client.get(
        f"/api/study-reports/{uuid.uuid4()}",
        headers={"X-User-Id": str(seeded.user.id)},
    )
    assert resp.status_code == 404


# ---------------------------------------------------------------------------
# Org isolation
# ---------------------------------------------------------------------------


async def _other_org_generator(db) -> User:
    other = User(
        id=uuid.uuid4(),
        org_id="other-org",
        email="generator@othercsr.com",
        role=UserRole.generator,
        hashed_password=hash_password("irrelevant"),
    )
    db.add(other)
    await db.commit()
    await db.refresh(other)
    return other


async def test_cross_org_cannot_read_csr(api_client, seeded, mock_anthropic, httpx_mock):
    protocol_id = await _create_protocol(api_client, seeded, httpx_mock)
    sr_id = await _create_study_report(api_client, seeded, protocol_id)
    await _upload(api_client, seeded, sr_id, _params_csv())

    other = await _other_org_generator(seeded.db)
    cross = await api_client.get(
        f"/api/study-reports/{sr_id}", headers={"X-User-Id": str(other.id)}
    )
    assert cross.status_code == 403, cross.text

    same = await api_client.get(
        f"/api/study-reports/{sr_id}", headers={"X-User-Id": str(seeded.user.id)}
    )
    assert same.status_code == 200


# ---------------------------------------------------------------------------
# Narrative generation (numbers-locked) + QC
# ---------------------------------------------------------------------------


async def test_generate_narrative_after_analysis(api_client, seeded, mock_anthropic, httpx_mock):
    """After analysis, generate the ICH E3 narrative (MOCK_AI → sync + credit-free).
    Numbers stay in results; sections carry prose only."""
    protocol_id = await _create_protocol(api_client, seeded, httpx_mock)
    sr_id = await _create_study_report(api_client, seeded, protocol_id)
    await _upload(api_client, seeded, sr_id, _params_csv())

    gen = await api_client.post(
        f"/api/study-reports/{sr_id}/generate",
        headers={"X-User-Id": str(seeded.user.id)},
    )
    assert gen.status_code == 200, gen.text
    body = gen.json()
    assert body["generation_status"] == "complete"
    # Every E3 section drafted; the results object is untouched.
    for key in ("synopsis", "bioequivalence_assessment", "safety_results", "discussion_conclusions"):
        assert body["sections"].get(key), f"missing section {key}"
    assert body["results"]["verdict"] == "bioequivalent"


async def test_generate_before_analysis_conflicts(api_client, seeded, mock_anthropic, httpx_mock):
    protocol_id = await _create_protocol(api_client, seeded, httpx_mock)
    sr_id = await _create_study_report(api_client, seeded, protocol_id)
    gen = await api_client.post(
        f"/api/study-reports/{sr_id}/generate",
        headers={"X-User-Id": str(seeded.user.id)},
    )
    assert gen.status_code == 409  # must analyse a dataset first


async def test_csr_qc(api_client, seeded, mock_anthropic, httpx_mock):
    protocol_id = await _create_protocol(api_client, seeded, httpx_mock)
    sr_id = await _create_study_report(api_client, seeded, protocol_id)
    await _upload(api_client, seeded, sr_id, _params_csv())

    qc = await api_client.get(
        f"/api/study-reports/{sr_id}/qc",
        headers={"X-User-Id": str(seeded.user.id)},
    )
    assert qc.status_code == 200, qc.text
    report = qc.json()
    assert report["submission_ready"] is True  # a clean bioequivalent result → no fails
    ids = {f["id"] for f in report["findings"]}
    assert "verdict" in ids and "coprimary_cmax" in ids and "coprimary_auc0inf" in ids


# ---------------------------------------------------------------------------
# ICH E3 export — the document renders the locked numbers verbatim
# ---------------------------------------------------------------------------


async def test_export_csr_docx_renders_locked_numbers(api_client, seeded, mock_anthropic, httpx_mock):
    import io as _io
    from docx import Document

    protocol_id = await _create_protocol(api_client, seeded, httpx_mock)
    sr_id = await _create_study_report(api_client, seeded, protocol_id)
    await _upload(api_client, seeded, sr_id, _params_csv())
    await api_client.post(f"/api/study-reports/{sr_id}/generate", headers={"X-User-Id": str(seeded.user.id)})

    # The exact point estimate the analysis computed (what MUST appear in the doc).
    sr_body = (await api_client.get(
        f"/api/study-reports/{sr_id}", headers={"X-User-Id": str(seeded.user.id)}
    )).json()
    pe_str = f"{sr_body['results']['be']['cmax']['point_estimate_pct']:.2f}"

    resp = await api_client.get(
        f"/api/study-reports/{sr_id}/export/docx",
        headers={"X-User-Id": str(seeded.user.id)},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml"
    )
    doc = Document(_io.BytesIO(resp.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                text += "\n" + cell.text
    assert "Clinical Study Report" in text
    assert pe_str in text, f"pivotal point estimate {pe_str} must render verbatim in the CSR"


async def test_export_csr_pdf_smoke(api_client, seeded, mock_anthropic, httpx_mock):
    protocol_id = await _create_protocol(api_client, seeded, httpx_mock)
    sr_id = await _create_study_report(api_client, seeded, protocol_id)
    await _upload(api_client, seeded, sr_id, _params_csv())

    resp = await api_client.get(
        f"/api/study-reports/{sr_id}/export/pdf",
        headers={"X-User-Id": str(seeded.user.id)},
    )
    assert resp.status_code == 200, resp.text
    assert resp.content[:4] == b"%PDF"


async def test_export_before_analysis_conflicts(api_client, seeded, mock_anthropic, httpx_mock):
    protocol_id = await _create_protocol(api_client, seeded, httpx_mock)
    sr_id = await _create_study_report(api_client, seeded, protocol_id)
    resp = await api_client.get(
        f"/api/study-reports/{sr_id}/export/docx",
        headers={"X-User-Id": str(seeded.user.id)},
    )
    assert resp.status_code == 409
