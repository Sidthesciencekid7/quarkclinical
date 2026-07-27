"""
Dataset ingestion for uploaded study results.

This module is the strict gate between a user-supplied spreadsheet and the
analysis engines (NCA / bioequivalence / safety). Its two jobs are:

1. Hand the user a *correct* template to fill in — as CSV or XLSX — so that a
   well-formed upload is easy to produce.
2. Parse an upload back into typed, validated Python rows, **rejecting bad data
   with a clear, per-row error message rather than silently mis-parsing it.**

Three dataset "kinds" are supported, each with a fixed column contract:

* ``"nca"``    — concentration-time samples (one row per subject/period/time).
* ``"params"`` — derived PK parameters (one row per subject/period).
* ``"ae"``     — adverse-event line listing.

Design notes
------------
* **Data-integrity first.** Every value is coerced with an explicit rule; a cell
  that does not satisfy its rule produces an error string like
  ``"row 4: 'treatment' must be T or R, got 'X'"``. Errors are *accumulated*
  across every field of every row (parsing does not stop at the first) and then
  raised together as a :class:`ValidationError`.
* **BLQ is missing, not zero.** In an NCA dataset a concentration that is blank
  or one of the tokens ``BLQ`` / ``<LLOQ`` / ``<LOQ`` (case-insensitive) parses
  to ``None`` — never ``0`` — so it cannot bias an AUC. If the caller passes an
  explicit ``lloq``, numeric concentrations strictly below it are treated the
  same way.
* **No side effects.** :func:`parse_dataset` returns parsed rows *in memory*
  together with a SHA-256 of the raw upload; it never writes to disk. The caller
  decides what (if anything) to retain.

Row-dict shapes returned by :func:`parse_dataset` (the orchestrator adapts these
to the NCA / BE engines):

* ``nca``    → ``{subject: str, sequence: str, period: int, treatment: 'T'|'R',
  time: float, concentration: float | None}``  (``None`` == below LOQ)
* ``params`` → ``{subject: str, sequence: str, period: int, treatment: 'T'|'R',
  cmax: float, auc0t: float, auc0inf: float, tmax: float | None,
  thalf: float | None}``
* ``ae``     → ``{subject: str, ae_term: str, severity: str, serious: bool,
  related: bool, start_day: int, outcome: str}``

Any additional (unrecognised) columns are preserved verbatim as stripped
strings under their own header name.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
from dataclasses import dataclass
from typing import Any

import openpyxl
from openpyxl.styles import Font

__all__ = [
    "ValidationError",
    "ParsedDataset",
    "template_columns",
    "build_template_csv",
    "build_template_xlsx",
    "parse_dataset",
    "REQUIRED_COLUMNS",
    "OPTIONAL_COLUMNS",
    "VALID_KINDS",
]

# ---------------------------------------------------------------------------
# Column contracts (lowercase, exact)
# ---------------------------------------------------------------------------

REQUIRED_COLUMNS: dict[str, list[str]] = {
    "nca": ["subject", "sequence", "period", "treatment", "time", "concentration"],
    "params": ["subject", "sequence", "period", "treatment", "cmax", "auc0t", "auc0inf"],
    "ae": ["subject", "ae_term", "severity", "serious", "related", "start_day", "outcome"],
}

OPTIONAL_COLUMNS: dict[str, list[str]] = {
    "nca": [],
    "params": ["tmax", "thalf"],
    "ae": [],
}

VALID_KINDS: tuple[str, ...] = tuple(REQUIRED_COLUMNS.keys())

# Tokens (case-insensitive) that mean "below limit of quantification" → None.
_BLQ_TOKENS = {"", "blq", "<lloq", "<loq"}
_TRUE_TOKENS = {"yes", "y", "true"}
_FALSE_TOKENS = {"no", "n", "false"}

# Cap the human-readable error message; ``.errors`` always holds the full list.
_MAX_ERRORS_SHOWN = 100

# Example data rows for the templates (native types so an XLSX round-trips
# cleanly). Values are obviously placeholders the user is told to delete.
_EXAMPLE_ROWS: dict[str, list[Any]] = {
    "nca": ["SUBJ-01", "TR", 1, "T", 0.5, 12.34],
    "params": ["SUBJ-01", "TR", 1, "T", 12.34, 45.6, 50.1, 1.5, 6.2],
    "ae": ["SUBJ-01", "Headache", "Mild", "no", "no", 2, "Recovered"],
}

# Comment block prepended to the NCA CSV template (parser skips ``#`` lines).
_NCA_TEMPLATE_COMMENTS = [
    "NCA concentration-time template. One row per (subject, period, time) sample.",
    "treatment must be T (test) or R (reference); period is a positive integer.",
    "Below-limit-of-quantification samples: leave the concentration cell blank OR",
    "enter BLQ / <LLOQ / <LOQ (case-insensitive). These parse to MISSING, never to",
    "0, so BLQ points do not bias the AUC. Delete this example row before uploading.",
]


# ---------------------------------------------------------------------------
# Public errors / result type
# ---------------------------------------------------------------------------


class ValidationError(Exception):
    """Raised when an upload fails validation.

    ``.errors`` holds the complete list of per-row/per-field error strings; the
    exception message renders up to ``_MAX_ERRORS_SHOWN`` of them followed by a
    "... and N more" note when the list is longer.
    """

    def __init__(self, errors: "str | list[str]") -> None:
        if isinstance(errors, str):
            errors = [errors]
        self.errors: list[str] = list(errors)
        n = len(self.errors)
        if n == 0:
            super().__init__("dataset validation failed")
            return
        shown = self.errors[:_MAX_ERRORS_SHOWN]
        body = "\n".join(f"  - {e}" for e in shown)
        if n > len(shown):
            body += f"\n  ... and {n - len(shown)} more"
        super().__init__(f"{n} validation error(s) in dataset:\n{body}")


@dataclass
class ParsedDataset:
    """A validated, typed upload held entirely in memory."""

    kind: str
    columns: list[str]  # header columns actually seen, in file order (lowercased)
    rows: list[dict]  # typed values; nca concentration == None means BLQ
    summary: dict  # n_rows, n_subjects, subjects, treatments, periods, sequences
    sha256: str  # hex digest of the raw uploaded bytes
    filename: str


class _FieldError(Exception):
    """Internal: one cell failed its rule. Carries a fully-formed message."""


# ---------------------------------------------------------------------------
# Template generators
# ---------------------------------------------------------------------------


def _check_kind(kind: str) -> None:
    if kind not in REQUIRED_COLUMNS:
        raise ValueError(
            f"unknown dataset kind {kind!r}; expected one of {', '.join(VALID_KINDS)}"
        )


def template_columns(kind: str) -> list[str]:
    """Full ordered column list for ``kind`` (required + optional)."""
    _check_kind(kind)
    return list(REQUIRED_COLUMNS[kind]) + list(OPTIONAL_COLUMNS[kind])


def build_template_csv(kind: str) -> bytes:
    """A ready-to-fill CSV template: header + one example row.

    For ``kind == "nca"`` a leading block of ``#`` comment lines documents the
    BLQ / LLOQ convention; :func:`parse_dataset` skips those comment lines.
    """
    _check_kind(kind)
    cols = template_columns(kind)
    buf = io.StringIO()
    if kind == "nca":
        for line in _NCA_TEMPLATE_COMMENTS:
            buf.write(f"# {line}\n")
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(cols)
    writer.writerow(_EXAMPLE_ROWS[kind])
    return buf.getvalue().encode("utf-8")


def build_template_xlsx(kind: str) -> bytes:
    """A ready-to-fill .xlsx template (bold header + one example row) as bytes."""
    _check_kind(kind)
    cols = template_columns(kind)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{kind}_template"
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(_EXAMPLE_ROWS[kind])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# Field parsers — each raises _FieldError with a complete "row N: ..." message
# ---------------------------------------------------------------------------


def _subject(raw: str, row_no: int) -> str:
    s = raw.strip()
    if not s:
        raise _FieldError(f"row {row_no}: 'subject' must be non-empty")
    return s


def _nonempty(raw: str, col: str, row_no: int) -> str:
    s = raw.strip()
    if not s:
        raise _FieldError(f"row {row_no}: '{col}' must be non-empty")
    return s


def _treatment(raw: str, row_no: int) -> str:
    s = raw.strip().upper()
    if s in ("T", "R"):
        return s
    raise _FieldError(f"row {row_no}: 'treatment' must be T or R, got {raw.strip()!r}")


def _positive_int(raw: str, col: str, row_no: int) -> int:
    s = raw.strip()
    try:
        f = float(s)
    except (TypeError, ValueError):
        raise _FieldError(f"row {row_no}: '{col}' must be a positive integer, got {s!r}")
    if not math.isfinite(f) or not f.is_integer() or f <= 0:
        raise _FieldError(f"row {row_no}: '{col}' must be a positive integer, got {s!r}")
    return int(f)


def _int(raw: str, col: str, row_no: int) -> int:
    s = raw.strip()
    try:
        f = float(s)
    except (TypeError, ValueError):
        raise _FieldError(f"row {row_no}: '{col}' must be an integer, got {s!r}")
    if not math.isfinite(f) or not f.is_integer():
        raise _FieldError(f"row {row_no}: '{col}' must be an integer, got {s!r}")
    return int(f)


def _nonneg_float(raw: str, col: str, row_no: int) -> float:
    s = raw.strip()
    if s == "":
        raise _FieldError(
            f"row {row_no}: '{col}' is required and must be a non-negative number"
        )
    try:
        f = float(s)
    except (TypeError, ValueError):
        raise _FieldError(
            f"row {row_no}: '{col}' must be a non-negative number, got {s!r}"
        )
    if not math.isfinite(f):
        raise _FieldError(f"row {row_no}: '{col}' must be a finite number, got {s!r}")
    if f < 0:
        raise _FieldError(f"row {row_no}: '{col}' must be >= 0, got {s!r}")
    return f


def _optional_nonneg_float(raw: str, col: str, row_no: int) -> "float | None":
    if raw.strip() == "":
        return None
    return _nonneg_float(raw, col, row_no)


def _bool(raw: str, col: str, row_no: int) -> bool:
    s = raw.strip().lower()
    if s in _TRUE_TOKENS:
        return True
    if s in _FALSE_TOKENS:
        return False
    raise _FieldError(
        f"row {row_no}: '{col}' must be one of yes/no/y/n/true/false, got {raw.strip()!r}"
    )


def _concentration(raw: str, row_no: int, lloq: "float | None") -> "float | None":
    s = raw.strip()
    if s.lower() in _BLQ_TOKENS:
        return None  # below limit of quantification → missing, NOT zero
    try:
        f = float(s)
    except (TypeError, ValueError):
        raise _FieldError(
            f"row {row_no}: 'concentration' must be a non-negative number or a BLQ "
            f"token (blank, BLQ, <LLOQ, <LOQ), got {raw.strip()!r}"
        )
    if not math.isfinite(f):
        raise _FieldError(
            f"row {row_no}: 'concentration' must be a finite number, got {s!r}"
        )
    if f < 0:
        raise _FieldError(f"row {row_no}: 'concentration' must be >= 0, got {s!r}")
    if lloq is not None and f < lloq:
        return None  # numerically below the caller-supplied LLOQ → treat as BLQ
    return f


# ---------------------------------------------------------------------------
# Low-level readers → list of (source_row_number, [raw string cells])
# ---------------------------------------------------------------------------


def _cell_to_str(value: Any) -> str:
    """Normalise an XLSX cell value to the string form the parsers expect."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        # Whole-number floats (e.g. openpyxl's 1.0 for a period) → "1", not "1.0".
        if math.isfinite(value) and value.is_integer():
            return str(int(value))
        return str(value)
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _is_skippable(cells: list[str]) -> bool:
    """True for a blank line or a leading-``#`` comment line."""
    if not cells or all(c.strip() == "" for c in cells):
        return True
    return cells[0].lstrip().startswith("#")


def _read_csv_records(file_bytes: bytes) -> list[tuple[int, list[str]]]:
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = file_bytes.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise ValidationError(
                f"file is not valid UTF-8 or Windows-1252 text: {exc}"
            )
    reader = csv.reader(io.StringIO(text))
    records: list[tuple[int, list[str]]] = []
    for cells in reader:
        row_no = reader.line_num  # physical line number (handles quoted newlines)
        if _is_skippable(cells):
            continue
        records.append((row_no, cells))
    return records


def _read_xlsx_records(file_bytes: bytes) -> list[tuple[int, list[str]]]:
    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(file_bytes), read_only=True, data_only=True
        )
    except Exception as exc:  # not a zip / corrupt / wrong format
        raise ValidationError(f"could not read .xlsx workbook: {exc}")
    try:
        ws = wb.worksheets[0]
        records: list[tuple[int, list[str]]] = []
        for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = [_cell_to_str(v) for v in row]
            if _is_skippable(cells):
                continue
            records.append((row_no, cells))
        return records
    finally:
        wb.close()


def _normalize_header(cells: list[str], row_no: int) -> tuple[list[str], dict[str, int]]:
    """Return (ordered lowercase names, name→column-index). Unnamed columns are
    ignored; duplicate names are a hard error."""
    names: list[str] = []
    index: dict[str, int] = {}
    dupes: list[str] = []
    for i, raw in enumerate(cells):
        name = raw.strip().lower()
        if name == "":
            continue  # e.g. a trailing comma / empty spacer column
        if name in index:
            dupes.append(name)
            continue
        index[name] = i
        names.append(name)
    if dupes:
        raise ValidationError(
            f"row {row_no}: duplicate column header(s): {', '.join(sorted(set(dupes)))}"
        )
    return names, index


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------


def parse_dataset(
    file_bytes: bytes,
    filename: str,
    kind: str,
    *,
    lloq: "float | None" = None,
) -> ParsedDataset:
    """Parse and strictly validate an uploaded dataset.

    Parameters
    ----------
    file_bytes:
        Raw bytes of the uploaded file.
    filename:
        Used only for format detection (``.csv`` / ``.xlsx``) and echoed back on
        the result.
    kind:
        One of :data:`VALID_KINDS`.
    lloq:
        Optional lower limit of quantification (NCA only). When given, numeric
        concentrations strictly below it are recorded as ``None`` (BLQ) in
        addition to the textual BLQ tokens.

    Returns
    -------
    ParsedDataset

    Raises
    ------
    ValueError
        If ``kind`` is not recognised (a programming error, not a data error).
    ValidationError
        If the format is unsupported, the header is missing required columns,
        or any data row fails its column rules (all row errors are accumulated).
    """
    _check_kind(kind)
    if not isinstance(file_bytes, (bytes, bytearray)):
        raise ValidationError("file_bytes must be raw bytes")
    raw = bytes(file_bytes)

    name = (filename or "").lower()
    if name.endswith(".csv"):
        records = _read_csv_records(raw)
    elif name.endswith(".xlsx"):
        records = _read_xlsx_records(raw)
    else:
        raise ValidationError(
            f"unsupported file type for {filename!r}; expected a .csv or .xlsx file"
        )

    if not records:
        raise ValidationError("file contains no header row or data rows")

    header_row_no, header_cells = records[0]
    names, index = _normalize_header(header_cells, header_row_no)

    required = REQUIRED_COLUMNS[kind]
    missing = [c for c in required if c not in index]
    if missing:
        raise ValidationError(
            f"header (row {header_row_no}) is missing required column(s) for kind "
            f"{kind!r}: {', '.join(missing)}. Found: {', '.join(names) or '(none)'}"
        )

    data_records = records[1:]
    if not data_records:
        raise ValidationError("no data rows found after the header row")

    known = set(required) | set(OPTIONAL_COLUMNS[kind])
    extra_cols = [c for c in names if c not in known]

    errors: list[str] = []
    rows: list[dict] = []
    seen_keys: dict[tuple, int] = {}  # (subject, period, time) → first row_no (nca)

    def cell_value(cells: list[str], col: str) -> str:
        idx = index.get(col)
        if idx is None or idx >= len(cells):
            return ""  # ragged row / column absent → empty, handled per-rule
        return cells[idx]

    for row_no, cells in data_records:
        row_errors: list[str] = []

        def grab(fn, *args):
            try:
                return fn(*args)
            except _FieldError as exc:
                row_errors.append(str(exc))
                return None

        rec: dict[str, Any] = {}

        if kind == "nca":
            rec["subject"] = grab(_subject, cell_value(cells, "subject"), row_no)
            rec["sequence"] = grab(_nonempty, cell_value(cells, "sequence"), "sequence", row_no)
            rec["period"] = grab(_positive_int, cell_value(cells, "period"), "period", row_no)
            rec["treatment"] = grab(_treatment, cell_value(cells, "treatment"), row_no)
            rec["time"] = grab(_nonneg_float, cell_value(cells, "time"), "time", row_no)
            rec["concentration"] = grab(
                _concentration, cell_value(cells, "concentration"), row_no, lloq
            )
            # Duplicate (subject, period, time) is a data-integrity failure: it
            # means the same sample was uploaded twice or two rows collide.
            if (
                not row_errors
                and rec["subject"] is not None
                and rec["period"] is not None
                and rec["time"] is not None
            ):
                key = (rec["subject"], rec["period"], rec["time"])
                if key in seen_keys:
                    row_errors.append(
                        f"row {row_no}: duplicate (subject, period, time) = "
                        f"({rec['subject']!r}, {rec['period']}, {rec['time']}); "
                        f"first seen at row {seen_keys[key]}"
                    )
                else:
                    seen_keys[key] = row_no

        elif kind == "params":
            rec["subject"] = grab(_subject, cell_value(cells, "subject"), row_no)
            rec["sequence"] = grab(_nonempty, cell_value(cells, "sequence"), "sequence", row_no)
            rec["period"] = grab(_positive_int, cell_value(cells, "period"), "period", row_no)
            rec["treatment"] = grab(_treatment, cell_value(cells, "treatment"), row_no)
            rec["cmax"] = grab(_nonneg_float, cell_value(cells, "cmax"), "cmax", row_no)
            rec["auc0t"] = grab(_nonneg_float, cell_value(cells, "auc0t"), "auc0t", row_no)
            rec["auc0inf"] = grab(_nonneg_float, cell_value(cells, "auc0inf"), "auc0inf", row_no)
            # Optional columns: always present in the dict (None when absent/blank).
            rec["tmax"] = (
                grab(_optional_nonneg_float, cell_value(cells, "tmax"), "tmax", row_no)
                if "tmax" in index
                else None
            )
            rec["thalf"] = (
                grab(_optional_nonneg_float, cell_value(cells, "thalf"), "thalf", row_no)
                if "thalf" in index
                else None
            )

        else:  # ae
            rec["subject"] = grab(_subject, cell_value(cells, "subject"), row_no)
            rec["ae_term"] = grab(_nonempty, cell_value(cells, "ae_term"), "ae_term", row_no)
            rec["severity"] = grab(_nonempty, cell_value(cells, "severity"), "severity", row_no)
            rec["serious"] = grab(_bool, cell_value(cells, "serious"), "serious", row_no)
            rec["related"] = grab(_bool, cell_value(cells, "related"), "related", row_no)
            rec["start_day"] = grab(_int, cell_value(cells, "start_day"), "start_day", row_no)
            rec["outcome"] = grab(_nonempty, cell_value(cells, "outcome"), "outcome", row_no)

        # Preserve any unrecognised columns verbatim (stripped strings).
        for col in extra_cols:
            rec[col] = cell_value(cells, col).strip()

        if row_errors:
            errors.extend(row_errors)
        else:
            rows.append(rec)

    if errors:
        raise ValidationError(errors)

    summary = _build_summary(kind, rows, lloq)
    sha256 = hashlib.sha256(raw).hexdigest()
    return ParsedDataset(
        kind=kind,
        columns=list(names),
        rows=rows,
        summary=summary,
        sha256=sha256,
        filename=filename,
    )


def _build_summary(kind: str, rows: list[dict], lloq: "float | None") -> dict:
    subjects = sorted({r["subject"] for r in rows if r.get("subject") is not None})
    treatments = sorted({r["treatment"] for r in rows if r.get("treatment") is not None})
    periods = sorted({r["period"] for r in rows if r.get("period") is not None})
    sequences = sorted({r["sequence"] for r in rows if r.get("sequence") is not None})
    summary: dict[str, Any] = {
        "n_rows": len(rows),
        "n_subjects": len(subjects),
        "subjects": subjects,
        "treatments": treatments,
        "periods": periods,
        "sequences": sequences,
    }
    if kind == "nca":
        summary["n_blq"] = sum(1 for r in rows if r.get("concentration") is None)
        summary["lloq"] = lloq
    return summary
